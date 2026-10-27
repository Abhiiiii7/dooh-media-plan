"""
Programmatic DOOH Media Plan & Reporting Model
-----------------------------------------------
Builds a sample DOOH media plan for a fictional grocery retailer across
3 metros (New York, Los Angeles, Chicago), then simulates the first two
weeks of delivery and produces:

  data/venues.csv          - venue inventory (simulated)
  data/media_plan.csv      - selected venues + planned spend/impressions
  data/delivery_daily.csv  - simulated daily delivery from publishers
  output/DOOH_Media_Plan.xlsx - Excel model with live formulas:
        Brief, Inventory, Media Plan, Delivery Data, Pacing, Recap Template

All venue, publisher and brand names are fictional. All numbers are simulated
for practice; nothing here is real campaign data.

Run:  python src/build_plan.py
"""
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA, OUT = ROOT / "data", ROOT / "output"
DATA.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)
rng = np.random.default_rng(42)

# ---------------------------------------------------------------- brief
BRIEF = {
    "Brand": "Brightside Grocers (fictional)",
    "Objective": "Drive store visits for spring produce launch",
    "Target audience": "Adults 25-54, primary grocery shoppers",
    "Flight start": "2026-03-02",
    "Flight end": "2026-03-29",
    "Flight days": 28,
    "Total budget ($)": 75000,
}
METRO_SPLIT = {"New York": 0.45, "Los Angeles": 0.35, "Chicago": 0.20}
DAYPARTS = ["Morning", "Midday", "Evening", "Weekend"]
DAYPART_WEIGHTS = {"Morning": 0.15, "Midday": 0.20, "Evening": 0.30, "Weekend": 0.35}
SELLABLE_SHARE = 0.40   # share of a venue's weekly impressions available to buy
BASE_TARGET_SHARE = 0.38  # share of general audience in target demo (index 100)
VENUE_CAP = 0.15        # max share of a metro budget on one venue
AS_OF_DAY = 14          # reporting date = day 14 of 28

# ------------------------------------------------------ venue inventory
VENUE_TYPES = {
    #  type        publisher           weekly impr      CPM         index        morning midday evening weekend
    "Transit":   ("MetroScreens",     (150_000, 400_000), (8, 12),  (95, 110),  (.35, .15, .35, .15)),
    "Gym":       ("FitView Network",  (20_000, 60_000),   (9, 13),  (105, 125), (.30, .20, .35, .15)),
    "Grocery":   ("AisleCast Media",  (40_000, 120_000),  (10, 15), (130, 160), (.15, .30, .30, .25)),
    "Office":    ("LobbyLink",        (30_000, 80_000),   (12, 18), (90, 105),  (.40, .40, .20, .00)),
    "Gas":       ("PumpPlay TV",      (50_000, 150_000),  (7, 10),  (105, 120), (.30, .20, .30, .20)),
    "Mall":      ("MallVision",       (80_000, 200_000),  (9, 13),  (115, 135), (.10, .30, .25, .35)),
}
METRO_PRICE = {"New York": 1.15, "Los Angeles": 1.05, "Chicago": 0.95}
METRO_CODE = {"New York": "NYC", "Los Angeles": "LAX", "Chicago": "CHI"}

rows = []
for metro in METRO_SPLIT:
    n = 0
    for vtype, (pub, impr, cpm, idx, mix) in VENUE_TYPES.items():
        for k in range(rng.integers(6, 10)):
            n += 1
            noise = rng.normal(0, 0.03, 4)
            m = np.clip(np.array(mix) + noise * (np.array(mix) > 0), 0, None)
            m = m / m.sum()
            rows.append({
                "venue_id": f"{METRO_CODE[metro]}-{n:03d}",
                "metro": metro,
                "publisher": pub,
                "venue_type": vtype,
                "screens": int(rng.integers(4, 40)),
                "weekly_impressions": int(round(rng.uniform(*impr), -3)),
                "cpm_floor": round(rng.uniform(*cpm) * METRO_PRICE[metro], 2),
                "audience_index": int(rng.integers(*idx)),
                **{f"share_{d.lower()}": round(float(s), 2) for d, s in zip(DAYPARTS, m)},
            })
venues = pd.DataFrame(rows)

# --------------------------------------------------------- venue scoring
w = np.array([DAYPART_WEIGHTS[d] for d in DAYPARTS])
shares = venues[[f"share_{d.lower()}" for d in DAYPARTS]].to_numpy()
venues["daypart_fit"] = (shares @ w) / 0.25          # 1.0 = average fit
venues["metro_avg_cpm"] = venues.groupby("metro")["cpm_floor"].transform("mean")
venues["score"] = (venues["audience_index"] / 100) * venues["daypart_fit"] / (
    venues["cpm_floor"] / venues["metro_avg_cpm"])

# ------------------------------------------------------- budget allocation
plan_rows = []
for metro, split in METRO_SPLIT.items():
    budget = BRIEF["Total budget ($)"] * split
    remaining = budget
    pool = venues[venues.metro == metro].sort_values("score", ascending=False)
    for _, v in pool.iterrows():
        if remaining < 500:
            break
        avail_impr = v.weekly_impressions * 4 * SELLABLE_SHARE
        max_spend = min(avail_impr * v.cpm_floor / 1000, budget * VENUE_CAP)
        spend = float(np.floor(min(max_spend, remaining) / 100) * 100)
        if spend <= 0:
            continue
        plan_rows.append({"venue_id": v.venue_id, "metro": metro, "publisher": v.publisher,
                          "venue_type": v.venue_type, "planned_spend": spend})
        remaining -= spend
plan = pd.DataFrame(plan_rows)
plan = plan.merge(venues[["venue_id", "cpm_floor", "audience_index"]], on="venue_id")
plan["planned_impressions"] = (plan.planned_spend / plan.cpm_floor * 1000).round()

# ---------------------------------------------------- simulated delivery
# Publisher delivery rates vs plan (what a real flight might look like):
# LobbyLink has screens offline; PumpPlay is inventory-constrained.
PUB_RATE = {"MetroScreens": 1.02, "FitView Network": 0.97, "AisleCast Media": 0.99,
            "LobbyLink": 0.82, "PumpPlay TV": 0.71, "MallVision": 1.04}
start = pd.Timestamp(BRIEF["Flight start"])
deliv = []
for _, p in plan.iterrows():
    daily_plan = p.planned_impressions / BRIEF["Flight days"]
    for d in range(AS_OF_DAY):
        date = start + pd.Timedelta(days=d)
        weekend = date.dayofweek >= 5
        wk_factor = 0.4 if (p.venue_type == "Office" and weekend) else (1.15 if weekend else 0.94)
        if p.venue_type == "Office" and not weekend:
            wk_factor = 1.24
        impr = daily_plan * PUB_RATE[p.publisher] * wk_factor * rng.normal(1, 0.06)
        impr = max(0, int(round(impr)))
        deliv.append({"date": date.date().isoformat(), "venue_id": p.venue_id, "metro": p.metro,
                      "publisher": p.publisher, "delivered_impressions": impr,
                      "delivered_spend": round(impr * p.cpm_floor / 1000, 2)})
delivery = pd.DataFrame(deliv)

venues.round(3).to_csv(DATA / "venues.csv", index=False)
plan.to_csv(DATA / "media_plan.csv", index=False)
delivery.to_csv(DATA / "delivery_daily.csv", index=False)

# ================================================================ EXCEL
F = "Arial"
BLUE, GREEN = Font(name=F, color="0000FF"), Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
BASE = Font(name=F)
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
MONEY, MONEY2 = "$#,##0;($#,##0);-", "$#,##0.00;($#,##0.00);-"
NUM, PCT, X = "#,##0;(#,##0);-", "0.0%;(0.0%);-", "0.00x"

wb = Workbook()


def header(ws, row, labels, widths=None):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font, cell.fill, cell.border = HEAD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for c, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = wd


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font() or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, color=c.font.color, italic=c.font.italic)


# ---------------- Brief
ws = wb.active
ws.title = "Brief"
ws["A1"] = "DOOH Media Plan - Campaign Brief"
ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "Blue = input you can change. Black = formula. All data simulated for practice."
ws["A2"].font = Font(name=F, italic=True, color="7F7F7F")
r = 4
brief_cells = {}
for k, v in BRIEF.items():
    ws.cell(row=r, column=1, value=k).font = BOLD
    c = ws.cell(row=r, column=2, value=v)
    c.font = BLUE
    brief_cells[k] = f"Brief!$B${r}"
    r += 1
ws[brief_cells["Total budget ($)"].split("!")[1].replace("$", "")].number_format = MONEY

r += 1
ws.cell(row=r, column=1, value="Metro budget split").font = BOLD
r += 1
metro_rows = {}
for m, s in METRO_SPLIT.items():
    ws.cell(row=r, column=1, value=m)
    ws.cell(row=r, column=2, value=s).font = BLUE
    ws.cell(row=r, column=2).number_format = PCT
    ws.cell(row=r, column=3, value=f"=B{r}*{brief_cells['Total budget ($)']}").number_format = MONEY
    metro_rows[m] = r
    r += 1
ws.cell(row=r, column=1, value="Check (must be 100%)").font = Font(name=F, italic=True)
first, last = min(metro_rows.values()), max(metro_rows.values())
ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{last})").number_format = PCT

r += 2
ws.cell(row=r, column=1, value="Daypart weights (what matters for this brand)").font = BOLD
r += 1
dp_rows = {}
for d in DAYPARTS:
    ws.cell(row=r, column=1, value=d)
    ws.cell(row=r, column=2, value=DAYPART_WEIGHTS[d]).font = BLUE
    ws.cell(row=r, column=2).number_format = PCT
    dp_rows[d] = r
    r += 1
ws.cell(row=dp_rows["Weekend"], column=3,
        value="Grocery trips skew to evenings and weekends, so those dayparts carry more weight.").font = Font(name=F, italic=True, color="7F7F7F")

r += 1
ws.cell(row=r, column=1, value="Planning assumptions").font = BOLD
r += 1
assume = {}
for label, val, fmt, note in [
    ("Sellable share of venue impressions", SELLABLE_SHARE, PCT, "Assumption: share of each venue's weekly impressions open to programmatic buys."),
    ("Target share at index 100", BASE_TARGET_SHARE, PCT, "Assumption: adults 25-54 grocery shoppers as a share of the general audience."),
    ("Max share of metro budget per venue", VENUE_CAP, PCT, "Keeps the plan spread across venues."),
    ("Reporting as-of (flight day)", AS_OF_DAY, "0", "Change this to see pacing at a different point in the flight."),
]:
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    c.font, c.number_format = BLUE, fmt
    if label.startswith("Reporting"):
        c.fill = YELLOW
    c.comment = Comment(note, "Model")
    assume[label] = f"Brief!$B${r}"
    r += 1
ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 16

DP_W = {d: f"Brief!$B${dp_rows[d]}" for d in DAYPARTS}
FLIGHT_DAYS = brief_cells["Flight days"]
AS_OF = assume["Reporting as-of (flight day)"]
TARGET_SHARE = assume["Target share at index 100"]

# ---------------- Inventory
wi = wb.create_sheet("Inventory")
cols = ["Venue ID", "Metro", "Publisher", "Venue Type", "Screens", "Weekly Impressions",
        "CPM Floor ($)", "Audience Index", "Morning %", "Midday %", "Evening %", "Weekend %",
        "Daypart Fit", "Metro Avg CPM ($)", "Venue Score", "Rank in Metro"]
header(wi, 1, cols, [11, 13, 18, 11, 9, 14, 11, 11, 10, 10, 10, 10, 11, 13, 11, 10])
n_inv = len(venues)
for i, v in enumerate(venues.itertuples(), start=2):
    vals = [v.venue_id, v.metro, v.publisher, v.venue_type, v.screens, v.weekly_impressions,
            v.cpm_floor, v.audience_index, v.share_morning, v.share_midday, v.share_evening, v.share_weekend]
    for c, val in enumerate(vals, 1):
        cell = wi.cell(row=i, column=c, value=val)
        cell.font = BASE
    wi.cell(row=i, column=13, value=(
        f"=(I{i}*{DP_W['Morning']}+J{i}*{DP_W['Midday']}+K{i}*{DP_W['Evening']}+L{i}*{DP_W['Weekend']})/0.25"))
    wi.cell(row=i, column=14, value=f"=AVERAGEIF($B$2:$B${n_inv+1},B{i},$G$2:$G${n_inv+1})")
    wi.cell(row=i, column=15, value=f"=(H{i}/100)*M{i}/(G{i}/N{i})")
    wi.cell(row=i, column=16, value=f"=COUNTIFS($B$2:$B${n_inv+1},B{i},$O$2:$O${n_inv+1},\">\"&O{i})+1")
    for c, fmt in [(5, NUM), (6, NUM), (7, MONEY2), (9, PCT), (10, PCT), (11, PCT), (12, PCT),
                   (13, "0.00"), (14, MONEY2), (15, "0.00"), (16, "0")]:
        wi.cell(row=i, column=c).number_format = fmt
wi.freeze_panes = "B2"
wi.auto_filter.ref = f"A1:P{n_inv+1}"
r = n_inv + 3
for line in ["Venue Score = (Audience Index / 100) x Daypart Fit / (CPM / Metro Avg CPM).",
             "Higher score = more of the right audience, at the right time, per dollar. Daypart Fit of 1.00 = average.",
             "Source: simulated inventory for practice (fictional publishers)."]:
    wi.cell(row=r, column=1, value=line).font = Font(name=F, italic=True, color="7F7F7F")
    r += 1

# ---------------- Media Plan
wp = wb.create_sheet("Media Plan")
cols = ["Venue ID", "Metro", "Publisher", "Venue Type", "Planned Spend ($)", "CPM ($)",
        "Planned Impressions", "Audience Index", "Est. Target Impressions", "Share of Metro Budget"]
header(wp, 1, cols, [11, 13, 18, 11, 15, 10, 16, 11, 17, 14])
INV = f"Inventory!$A$2:$A${n_inv+1}"
n_plan = len(plan)
for i, p in enumerate(plan.itertuples(), start=2):
    wp.cell(row=i, column=1, value=p.venue_id).font = BASE
    wp.cell(row=i, column=2, value=f"=INDEX(Inventory!$B$2:$B${n_inv+1},MATCH(A{i},{INV},0))").font = GREEN
    wp.cell(row=i, column=3, value=f"=INDEX(Inventory!$C$2:$C${n_inv+1},MATCH(A{i},{INV},0))").font = GREEN
    wp.cell(row=i, column=4, value=f"=INDEX(Inventory!$D$2:$D${n_inv+1},MATCH(A{i},{INV},0))").font = GREEN
    c = wp.cell(row=i, column=5, value=p.planned_spend)
    c.font = BLUE
    wp.cell(row=i, column=6, value=f"=INDEX(Inventory!$G$2:$G${n_inv+1},MATCH(A{i},{INV},0))").font = GREEN
    wp.cell(row=i, column=7, value=f"=IFERROR(E{i}/F{i}*1000,0)")
    wp.cell(row=i, column=8, value=f"=INDEX(Inventory!$H$2:$H${n_inv+1},MATCH(A{i},{INV},0))").font = GREEN
    wp.cell(row=i, column=9, value=f"=G{i}*{TARGET_SHARE}*H{i}/100")
    wp.cell(row=i, column=10, value=f"=IFERROR(E{i}/SUMIF($B$2:$B${n_plan+1},B{i},$E$2:$E${n_plan+1}),0)")
    for c, fmt in [(5, MONEY), (6, MONEY2), (7, NUM), (8, "0"), (9, NUM), (10, PCT)]:
        wp.cell(row=i, column=c).number_format = fmt
t = n_plan + 2
wp.cell(row=t, column=1, value="TOTAL").font = BOLD
for c, f_ in [(5, f"=SUM(E2:E{n_plan+1})"), (7, f"=SUM(G2:G{n_plan+1})"), (9, f"=SUM(I2:I{n_plan+1})")]:
    cell = wp.cell(row=t, column=c, value=f_)
    cell.font, cell.number_format = BOLD, (MONEY if c == 5 else NUM)
wp.cell(row=t, column=6, value=f"=IFERROR(E{t}/G{t}*1000,0)").number_format = MONEY2
for c in range(1, 11):
    wp.cell(row=t, column=c).fill = TOTAL_FILL
PLAN_TOTAL_ROW = t

r = t + 2
wp.cell(row=r, column=1, value="By metro").font = BOLD
r += 1
for c, lab in enumerate(["Metro", "Budget ($)", "Planned Spend ($)", "Difference ($)",
                         "Planned Impressions", "Blended CPM ($)", "Venues"], 1):
    cell = wp.cell(row=r, column=c, value=lab)
    cell.font, cell.fill = HEAD, HEAD_FILL
r += 1
PR = f"$B$2:$B${n_plan+1}"
for m in METRO_SPLIT:
    wp.cell(row=r, column=1, value=m)
    wp.cell(row=r, column=2, value=f"=Brief!$C${metro_rows[m]}").font = GREEN
    wp.cell(row=r, column=3, value=f"=SUMIF({PR},A{r},$E$2:$E${n_plan+1})")
    wp.cell(row=r, column=4, value=f"=B{r}-C{r}")
    wp.cell(row=r, column=5, value=f"=SUMIF({PR},A{r},$G$2:$G${n_plan+1})")
    wp.cell(row=r, column=6, value=f"=IFERROR(C{r}/E{r}*1000,0)")
    wp.cell(row=r, column=7, value=f"=COUNTIF({PR},A{r})")
    for c, fmt in [(2, MONEY), (3, MONEY), (4, MONEY), (5, NUM), (6, MONEY2), (7, "0")]:
        wp.cell(row=r, column=c).number_format = fmt
    r += 1
wp.cell(row=r + 1, column=1,
        value="Planned Spend is the input (blue). Venues were picked by Venue Score until each metro budget was used, "
              "capped by sellable impressions and the per-venue limit on the Brief.").font = Font(name=F, italic=True, color="7F7F7F")
wp.freeze_panes = "B2"

# ---------------- Delivery Data
wd = wb.create_sheet("Delivery Data")
cols = ["Date", "Flight Day", "Venue ID", "Metro", "Publisher", "Delivered Impressions", "Delivered Spend ($)"]
header(wd, 1, cols, [12, 10, 11, 13, 18, 18, 16])
n_del = len(delivery)
for i, dlv in enumerate(delivery.itertuples(), start=2):
    day = (pd.Timestamp(dlv.date) - start).days + 1
    for c, val in enumerate([dlv.date, day, dlv.venue_id, dlv.metro, dlv.publisher,
                             dlv.delivered_impressions, dlv.delivered_spend], 1):
        wd.cell(row=i, column=c, value=val).font = BASE
    wd.cell(row=i, column=6).number_format = NUM
    wd.cell(row=i, column=7).number_format = MONEY2
wd.freeze_panes = "A2"
wd.auto_filter.ref = f"A1:G{n_del+1}"
DD = lambda col: f"'Delivery Data'!${col}$2:${col}${n_del+1}"

# ---------------- Pacing
wpc = wb.create_sheet("Pacing")
wpc["A1"] = "Pacing & Delivery - by Publisher and Metro"
wpc["A1"].font = Font(name=F, bold=True, size=14)
wpc["A2"] = "=\"As of flight day \"&" + AS_OF + "&\" of \"&" + FLIGHT_DAYS
wpc["A2"].font = Font(name=F, italic=True)
wpc["D2"] = "% of flight elapsed"
wpc["E2"] = f"={AS_OF}/{FLIGHT_DAYS}"
wpc["E2"].number_format = PCT
wpc["E2"].font = BOLD
wpc["G2"] = "Behind if pace <"
wpc["H2"] = 0.9
wpc["H2"].font, wpc["H2"].number_format = BLUE, PCT
wpc["I2"] = "Ahead if pace >"
wpc["J2"] = 1.1
wpc["J2"].font, wpc["J2"].number_format = BLUE, PCT

pcols = ["Publisher", "Planned Impressions", "Expected to Date", "Delivered to Date (as-of)",
         "% of Plan Delivered", "Pace Index", "Projected Final Impressions", "Projected Shortfall",
         "Delivered Spend ($)", "Status"]
header(wpc, 4, pcols, [18, 16, 15, 17, 13, 11, 18, 16, 15, 12])
pubs = list(PUB_RATE.keys())
MP = lambda col: f"'Media Plan'!${col}$2:${col}${n_plan+1}"


def pacing_block(start_row, keys, plan_key_col, del_key_col):
    rr = start_row
    for k in keys:
        wpc.cell(row=rr, column=1, value=k)
        wpc.cell(row=rr, column=2, value=f"=SUMIF({MP(plan_key_col)},A{rr},{MP('G')})")
        wpc.cell(row=rr, column=3, value=f"=B{rr}*$E$2")
        wpc.cell(row=rr, column=4, value=f"=SUMIFS({DD('F')},{DD(del_key_col)},A{rr},{DD('B')},\"<=\"&{AS_OF})")
        wpc.cell(row=rr, column=5, value=f"=IFERROR(D{rr}/B{rr},0)")
        wpc.cell(row=rr, column=6, value=f"=IFERROR(D{rr}/C{rr},0)")
        wpc.cell(row=rr, column=7, value=f"=IFERROR(D{rr}/$E$2,0)")
        wpc.cell(row=rr, column=8, value=f"=MAX(0,B{rr}-G{rr})")
        wpc.cell(row=rr, column=9, value=f"=SUMIFS({DD('G')},{DD(del_key_col)},A{rr},{DD('B')},\"<=\"&{AS_OF})")
        wpc.cell(row=rr, column=10, value=f"=IF(B{rr}=0,\"Not in plan\",IF(F{rr}<$H$2,\"Behind\",IF(F{rr}>$J$2,\"Ahead\",\"On track\")))")
        for c, fmt in [(2, NUM), (3, NUM), (4, NUM), (5, PCT), (6, X), (7, NUM), (8, NUM), (9, MONEY)]:
            wpc.cell(row=rr, column=c).number_format = fmt
        rr += 1
    wpc.cell(row=rr, column=1, value="TOTAL").font = BOLD
    for c in [2, 3, 4, 7, 8, 9]:
        L = get_column_letter(c)
        wpc.cell(row=rr, column=c, value=f"=SUM({L}{start_row}:{L}{rr-1})").number_format = MONEY if c == 9 else NUM
    wpc.cell(row=rr, column=5, value=f"=IFERROR(D{rr}/B{rr},0)").number_format = PCT
    wpc.cell(row=rr, column=6, value=f"=IFERROR(D{rr}/C{rr},0)").number_format = X
    # net shortfall for the whole campaign (over-delivery on one line offsets another)
    wpc.cell(row=rr, column=8, value=f"=MAX(0,B{rr}-G{rr})").number_format = NUM
    wpc.cell(row=rr, column=10, value=f"=IF(F{rr}<$H$2,\"Behind\",IF(F{rr}>$J$2,\"Ahead\",\"On track\"))")
    for c in range(1, 11):
        wpc.cell(row=rr, column=c).fill = TOTAL_FILL
        wpc.cell(row=rr, column=c).font = BOLD
    return rr


pub_total = pacing_block(5, pubs, "C", "E")
mrow = pub_total + 2
header(wpc, mrow, ["Metro"] + pcols[1:])
metro_total = pacing_block(mrow + 1, list(METRO_SPLIT), "B", "D")
note = metro_total + 2
for line in ["Pace Index = Delivered to Date / Expected to Date. 1.00x = exactly on plan.",
             "Projected Final assumes the current daily rate holds for the rest of the flight.",
             "Change the as-of day on the Brief sheet to re-run pacing at another point in the flight."]:
    wpc.cell(row=note, column=1, value=line).font = Font(name=F, italic=True, color="7F7F7F")
    note += 1
PUB_FIRST, PUB_LAST, PUB_TOTAL = 5, pub_total - 1, pub_total

# ---------------- Recap Template
wr = wb.create_sheet("Recap Template")
wr["A1"] = "Campaign Recap - Brightside Grocers Spring Produce (DOOH)"
wr["A1"].font = Font(name=F, bold=True, size=14)
wr["A2"] = "Headline numbers fill in automatically. Fill the yellow cells for the client-facing story."
wr["A2"].font = Font(name=F, italic=True, color="7F7F7F")
kpis = [
    ("Planned impressions", f"='Pacing'!B{PUB_TOTAL}", NUM),
    ("Delivered impressions (to date)", f"='Pacing'!D{PUB_TOTAL}", NUM),
    ("% of plan delivered", f"='Pacing'!E{PUB_TOTAL}", PCT),
    ("Overall pace index", f"='Pacing'!F{PUB_TOTAL}", X),
    ("Delivered spend ($)", f"='Pacing'!I{PUB_TOTAL}", MONEY),
    ("Effective CPM ($)", f"=IFERROR(B8/B5*1000,0)", MONEY2),
    ("Est. target impressions delivered", f"=B5*'Media Plan'!I{PLAN_TOTAL_ROW}/'Media Plan'!G{PLAN_TOTAL_ROW}", NUM),
    ("Publishers behind pace", f"=COUNTIF('Pacing'!J{PUB_FIRST}:J{PUB_LAST},\"Behind\")", "0"),
]
r = 4
for lab, f_, fmt in kpis:
    wr.cell(row=r, column=1, value=lab).font = BOLD
    c = wr.cell(row=r, column=2, value=f_)
    c.number_format, c.font = fmt, GREEN
    r += 1
r += 1
sections = [
    ("What worked", "Grocery-venue screens delivered on plan at the strongest audience index; weekend dayparts carried the reach."),
    ("What under-delivered and why", "PumpPlay TV ran ~30% behind due to limited weekday inventory; flagged on day 7."),
    ("What we changed mid-flight", "Shifted unspent PumpPlay budget to MallVision weekend inventory in Los Angeles."),
    ("Recommendation for next flight", "Raise grocery and mall share; set make-good terms with gas-station publishers up front."),
]
wr.cell(row=r, column=1, value="Story for the client").font = BOLD
wr.cell(row=r, column=2, value="Your notes (example row in grey - replace it)").font = BOLD
r += 1
for lab, example in sections:
    wr.cell(row=r, column=1, value=lab).font = BOLD
    c = wr.cell(row=r, column=2, value=example)
    c.fill, c.font = YELLOW, Font(name=F, italic=True, color="7F7F7F")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    wr.row_dimensions[r].height = 45
    r += 1
wr.column_dimensions["A"].width = 34
wr.column_dimensions["B"].width = 80

for s in wb.worksheets:
    style_all(s)
wb.save(OUT / "DOOH_Media_Plan.xlsx")
print(f"venues={len(venues)} plan_lines={len(plan)} delivery_rows={len(delivery)}")
print(plan.groupby("metro")[["planned_spend", "planned_impressions"]].sum())
